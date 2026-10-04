"""Таблица для вычитки носителем языка: все формы ~200 имён и слов с колонками для правок.

python review/build_review_table.py → review/qazaq_inflector_review.xlsx
"""
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from qazaq_inflector import QazaqNameInflector, __version__  # noqa: E402

GROUPS = {
    'Ер есімі / мужское имя': [
        'Нұрлан', 'Абай', 'Ерлан', 'Ержан', 'Мейрам', 'Айбаз', 'Сұлтанбек', 'Бақыт', 'Дәулет', 'Асқар', 'Мұрат',
        'Ғалым', 'Төлеген', 'Қуаныш', 'Еркебұлан', 'Олжас', 'Нұрсұлтан', 'Серік', 'Ілияс', 'Өмірзақ', 'Айдос',
        'Ринат', 'Азамат', 'Ахмет', 'Бауыржан', 'Дархан', 'Ербол', 'Жандос', 'Қайрат', 'Мақсат', 'Нұржан', 'Самат',
        'Талғат', 'Ұлан', 'Шыңғыс', 'Әділет', 'Бекзат', 'Ғабит', 'Ділмұрат', 'Еділ', 'Жасұлан', 'Кенжебек',
        'Мағжан', 'Нұрболат', 'Темірлан', 'Санжар', 'Әлібек', 'Мирас', 'Арман', 'Тілек',
    ],
    'Әйел есімі / женское имя': [
        'Арна', 'Сәуле', 'Айгүл', 'Ләззат', 'Мәдина', 'Гүлнар', 'Жанар', 'Ботагөз', 'Динара', 'Әсел', 'Ақбота',
        'Шолпан', 'Аружан', 'Үміт', 'Іңкәр', 'Ұлжан', 'Айгерім', 'Балжан', 'Гүлмира', 'Дана', 'Жұлдыз', 'Камила',
        'Ләйлә', 'Мөлдір', 'Назерке', 'Перизат', 'Салтанат', 'Томирис', 'Әйгерім', 'Гаухар', 'Ақмарал', 'Меруерт',
        'Алтынай', 'Еркежан', 'Зере', 'Күнсұлу', 'Нәзира', 'Раушан', 'Сымбат', 'Хадиша',
    ],
    'Кірме есім / заимствованное имя': [
        'Әли', 'Мұхаммед', 'Жүсіп', 'Ахмед', 'Тимур', 'Ибраһим', 'Мария', 'Анастасия', 'Дмитрий', 'Виктор',
        'Юрий', 'Ольга', 'Евгений', 'Наталья', 'Сергей', 'Карина', 'Эльвира', 'Рустем', 'Фатима', 'Хасан',
    ],
    'Тегі (ер) / фамилия (муж.)': [
        'Құнанбаев', 'Ахметов', 'Әлиев', 'Иванов', 'Назарбаев', 'Тоқаев', 'Сәтбаев', 'Байжанов', 'Жұмабеков',
        'Кенесарин', 'Әуезов', 'Мұқанов', 'Есенов', 'Өтебаев', 'Ғабдуллин', 'Ыбыраев', 'Сейфуллин', 'Шәріпов',
        'Ермеков', 'Дүйсенов',
    ],
    'Тегі (әйел) / фамилия (жен.)': [
        'Назарбаева', 'Ахметова', 'Әлиева', 'Сейітова', 'Тоқаева', 'Сәтбаева', 'Жұмабекова', 'Әуезова',
        'Мұқанова', 'Есенова', 'Өтебаева', 'Иванова', 'Ермекова', 'Дүйсенова', 'Ыбыраева',
    ],
    'Әкесінің аты / отчество': ['Байтұрсынұлы', 'Нұрланқызы', 'Абайұлы', 'Серікқызы'],
    'ФИО / толық аты-жөні': [
        'Абай Құнанбаев', 'Ахмет Байтұрсынұлы', 'Әлия Нұрланқызы', 'Мұхтар Әуезов', 'Сәуле Ахметова',
        'Гүлнар-Баян', 'Әли-Хан',
    ],
    'Жалпы зат есім / нарицательное': [
        'кітап', 'әке', 'ана', 'бала', 'үй', 'дос', 'мұғалім', 'дәрігер', 'оқушы', 'студент', 'қала', 'көл', 'ақ',
        'жер', 'тау', 'су', 'ауыл', 'мектеп', 'бағ', 'қыз', 'жас', 'құс', 'сөз', 'гүл', 'ит', 'күз', 'бет', 'ағаш',
        'терезе', 'есік',
    ],
    'Есімдік / местоимение': ['мен', 'сен', 'сіз', 'ол', 'біз', 'сендер', 'сіздер', 'олар'],
}

CASE_LABELS = ['Атау', 'Ілік', 'Барыс', 'Табыс', 'Жатыс', 'Шығыс', 'Көмектес']
PERSONS = [('1sg', 'менің'), ('2sg', 'сенің'), ('2sg_formal', 'сіздің'), ('1pl', 'біздің'),
           ('2pl', 'сендердің'), ('2pl_formal', 'сіздердің'), ('3', 'оның')]
PRED_PERSONS = [('1sg', 'мен …'), ('2sg', 'сен …'), ('2sg_formal', 'сіз …'), ('1pl', 'біз …'),
                ('2pl', 'сендер …'), ('2pl_formal', 'сіздер …')]

FONT = 'Arial'
HEADER_FILL = PatternFill('solid', start_color='1F3864')
GROUP_FILL = PatternFill('solid', start_color='D9E1F2')
INPUT_FILL = PatternFill('solid', start_color='FFF2CC')
EXAMPLE_FILL = PatternFill('solid', start_color='E2EFDA')
ERROR_FILL = PatternFill('solid', start_color='F8CBAD')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
REVIEW_HEADERS = ['Дұрыс па? / Верно?', 'Дұрыс нұсқасы / Правильно', 'Ескерту / Комментарий']

inflector = QazaqNameInflector()


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 32


def build_sheet(wb, title, headers, rows_for, widths, example):
    """rows_for(word) → список форм; example — (слово, ✓/✗, правка, комментарий) для образца."""
    ws = wb.create_sheet(title)
    all_headers = ['Сөз / Слово'] + headers + REVIEW_HEADERS
    ws.append(all_headers)
    style_header(ws, 1, len(all_headers))
    review_col = len(headers) + 2
    validation = DataValidation(type='list', formula1='"✓,✗"', allow_blank=True)
    ws.add_data_validation(validation)

    word, mark, fix, note = example
    ws.append([word] + rows_for(word) + [mark, fix, note])
    for col in range(1, len(all_headers) + 1):
        ws.cell(row=2, column=col).fill = EXAMPLE_FILL
        ws.cell(row=2, column=col).font = Font(name=FONT, italic=True)
        ws.cell(row=2, column=col).border = BORDER
    ws.cell(row=2, column=1).comment = Comment(
        'Үлгі жол — толтыру үлгісі. Образец заполнения, не проверять.', 'qazaq_inflector')

    for group, words in GROUPS.items():
        group_rows = [(w, rows_for(w)) for w in words]
        group_rows = [(w, forms) for w, forms in group_rows if forms is not None]
        if not group_rows:
            continue
        ws.append([group])
        r = ws.max_row
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(all_headers))
        ws.cell(row=r, column=1).font = Font(name=FONT, bold=True)
        ws.cell(row=r, column=1).fill = GROUP_FILL
        for w, forms in group_rows:
            ws.append([w] + forms + ['', '', ''])
            r = ws.max_row
            for col in range(1, len(all_headers) + 1):
                cell = ws.cell(row=r, column=col)
                cell.font = Font(name=FONT, bold=(col == 1))
                cell.border = BORDER
                if col >= review_col:
                    cell.fill = INPUT_FILL
            validation.add(ws.cell(row=r, column=review_col))

    last = ws.max_row
    rc = ws.cell(row=1, column=review_col).column_letter
    ws.conditional_formatting.add(
        f'A3:{ws.cell(row=1, column=len(all_headers)).column_letter}{last}',
        FormulaRule(formula=[f'${rc}3="✗"'], fill=ERROR_FILL),
    )
    for idx, width in enumerate([18] + widths + [14, 24, 30], start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = 'B2'
    return ws, rc, last


def cases_row(word):
    plural = inflector.pluralize(word)
    return [inflector.inflect(word, c) for c in QazaqNameInflector.CASES] + \
        [inflector.inflect(plural, c) for c in QazaqNameInflector.CASES]


def possessive_row(word):
    if word in GROUPS['Есімдік / местоимение'] or word in GROUPS['Әкесінің аты / отчество']:
        return None
    forms = [inflector.possessive(word, p) for p, _ in PERSONS]
    forms += [inflector.possessive(word, '3', c) for c in QazaqNameInflector.CASES[1:]]
    forms += [inflector.possessive(word, '1sg', 'dative'), inflector.possessive(word, '1sg', plural=True)]
    return forms


def predicate_row(word):
    if word in GROUPS['Есімдік / местоимение'] or ' ' in word:
        return None
    return [inflector.predicate(word, p) for p, _ in PRED_PERSONS]


wb = Workbook()
guide = wb.active
guide.title = 'Нұсқаулық'

case_headers = [f'{c} (жекеше)' for c in CASE_LABELS] + [f'{c} (көпше)' for c in CASE_LABELS]
poss_headers = [f'{p} ({owner})' for p, owner in PERSONS] + \
    [f'оның + {c.lower()}' for c in CASE_LABELS[1:]] + ['менің + барыс', 'менің + көпше']
pred_headers = [f'{p} ({pr})' for p, pr in PRED_PERSONS]

sheets = [
    build_sheet(wb, 'Септік', case_headers, cases_row, [16] * 14,
                ('Нұрлан', '✗', 'Шығыс: Нұрланнан', 'Үлгі: қате болса ✗ қойып, дұрысын жазыңыз')),
    build_sheet(wb, 'Тәуелдік', poss_headers, possessive_row, [16] * 15,
                ('Арна', '✓', '', 'Үлгі: бәрі дұрыс болса ✓')),
    build_sheet(wb, 'Жіктік', pred_headers, predicate_row, [18] * 6,
                ('студент', '✓', '', 'Үлгі')),
]

rows = [
    ('qazaq_inflector — тексеру кестесі / таблица для вычитки', True),
    (f'Нұсқа / версия библиотеки: {__version__}', False),
    ('', False),
    ('Қазақша', True),
    ('Кестеде кітапхана автоматты түрде жасаған сөз формалары берілген. Әр жолды тексеріңіз:', False),
    ('• бәрі дұрыс болса — «Дұрыс па?» бағанына ✓ қойыңыз;', False),
    ('• қате болса — ✗ қойып, «Дұрыс нұсқасы» бағанына қай формасы қате екенін және дұрысын жазыңыз;', False),
    ('• күмәнді болса — «Ескерту» бағанына жазыңыз (мысалы, «екі нұсқа да қолданылады»).', False),
    ('Тек сары ұяшықтарды толтырыңыз. Жасыл жол — толтыру үлгісі, оны тексермеңіз.', False),
    ('', False),
    ('По-русски', True),
    ('В таблице — формы, которые библиотека строит автоматически. Проверьте каждую строку:', False),
    ('• всё верно — поставьте ✓ в колонке «Верно?»;', False),
    ('• есть ошибка — поставьте ✗ и в колонке «Правильно» укажите, какая форма неверна и как правильно;', False),
    ('• сомневаетесь — напишите в «Комментарий» (например, «допустимы оба варианта»).', False),
    ('Заполняйте только жёлтые ячейки. Зелёная строка — образец заполнения, её проверять не нужно.', False),
    ('', False),
    ('Парақтар / Листы', True),
    ('Септік — 7 септік, жекеше және көпше / падежи в единственном и множественном числе', False),
    ('Тәуелдік — тәуелдік жалғау, барлық жақ және септікпен / притяжательные формы', False),
    ('Жіктік — жіктік жалғау: студентпін, Нұрлансың / личные окончания сказуемого', False),
    ('', False),
    ('Барысы / Прогресс', True),
]
for text, bold in rows:
    guide.append([text])
    guide.cell(row=guide.max_row, column=1).font = Font(name=FONT, bold=bold, size=14 if guide.max_row == 1 else 11)

guide.append(['Парақ / Лист', 'Барлығы / Всего', '✓', '✗', 'Тексерілмеген / Осталось'])
style_header(guide, guide.max_row, 5)
for ws, rc, last in sheets:
    guide.append([ws.title])
    r = guide.max_row
    rng = f"'{ws.title}'!{rc}3:{rc}{last}"
    words = f"'{ws.title}'!B3:B{last}"
    guide.cell(row=r, column=2, value=f'=COUNTA({words})')
    guide.cell(row=r, column=3, value=f'=COUNTIF({rng},"✓")')
    guide.cell(row=r, column=4, value=f'=COUNTIF({rng},"✗")')
    guide.cell(row=r, column=5, value=f'=B{r}-C{r}-D{r}')
    for col in range(1, 6):
        guide.cell(row=r, column=col).font = Font(name=FONT, color='008000' if col in (2, 3, 4) else '000000')
        guide.cell(row=r, column=col).border = BORDER
guide.column_dimensions['A'].width = 24
for col in 'BCDE':
    guide.column_dimensions[col].width = 18
for r in range(1, 23):
    guide.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    guide.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')

out = pathlib.Path(__file__).with_name('qazaq_inflector_review.xlsx')
wb.save(out)
print(out, {ws.title: last - 2 for ws, _, last in sheets})
